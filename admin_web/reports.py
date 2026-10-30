from datetime import datetime, timezone
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font

import settings
import stats


def _style_header(ws, ncols: int):
    bold = Font(bold=True)
    for col in range(1, ncols + 1):
        ws.cell(row=1, column=col).font = bold
    widths = [22, 20, 16, 14, 14, 34, 16, 12, 22, 22, 16, 24]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w


def build_report() -> tuple[bytes, str]:
    overview = stats.get_overview()
    clients = stats.get_clients()
    sales = stats.get_sales(30)
    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric", "Value"])
    for key, value in overview.items():
        if isinstance(value, dict):
            for sub, num in value.items():
                ws.append([f"{key}.{sub}", num])
        else:
            ws.append([key, value])
    _style_header(ws, 2)

    ws2 = wb.create_sheet("Clients")
    cols = ["user_id", "username", "full_name", "plan", "status", "method", "tx_id",
            "amount_usd", "hired_at", "expires_at", "verify_status", "demo"]
    ws2.append(cols)
    for c in clients:
        ws2.append([c.get(k) for k in cols])
    _style_header(ws2, len(cols))

    ws3 = wb.create_sheet("Sales30d")
    ws3.append(["date", "revenue_usd", "sales"])
    for s in sales:
        ws3.append([s["date"], s["revenue"], s["sales"]])
    _style_header(ws3, 3)

    buf = BytesIO()
    wb.save(buf)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d")
    return buf.getvalue(), f"hire-report-{stamp}.xlsx"
